"""Build a visual Excel QA/QC dashboard for the Client XYZ delivery."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NAVY = "1F4E78"
BLUE = "5B9BD5"
GREEN = "70AD47"
AMBER = "FFC000"
RED = "C00000"
GRAY = "7F8C8D"
WHITE = "FFFFFF"
DARK = "1F1F1F"

STATUS_COLORS = {
    "READY_TO_UPDATE": GREEN,
    "NO_CHANGE": BLUE,
    "REJECTED": RED,
    "UNMATCHED": AMBER,
}


def _safe(value: Any) -> Any:
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value


def _write_dataframe(ws, df: pd.DataFrame) -> None:
    headers = list(df.columns)
    ws.append(headers)
    for row in df.itertuples(index=False, name=None):
        ws.append([_safe(v) for v in row])

    header_fill = PatternFill("solid", fgColor=NAVY)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        max_len = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells[: min(len(column_cells), 250)]
        )
        ws.column_dimensions[letter].width = min(max(max_len + 2, 11), 32)


def _card(ws, start_col: int, title: str, value: Any, fill_color: str) -> None:
    end_col = start_col + 1
    ws.merge_cells(start_row=5, start_column=start_col, end_row=5, end_column=end_col)
    ws.merge_cells(start_row=6, start_column=start_col, end_row=7, end_column=end_col)

    title_cell = ws.cell(5, start_col)
    value_cell = ws.cell(6, start_col)
    title_cell.value = title
    value_cell.value = value

    for row in range(5, 8):
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row, col)
            cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(
                left=Side(style="thin", color=WHITE),
                right=Side(style="thin", color=WHITE),
                top=Side(style="thin", color=WHITE),
                bottom=Side(style="thin", color=WHITE),
            )

    title_cell.font = Font(color=WHITE, bold=True, size=10)
    value_cell.font = Font(color=WHITE, bold=True, size=20)


def generate_dashboard(
    output_path: str | Path,
    summary: dict[str, Any],
    reports: dict[str, pd.DataFrame],
    classified: pd.DataFrame,
) -> Path:
    """Create a corporate-style Excel QA/QC workbook with a dashboard."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:L2")
    ws["A1"] = "CLIENT XYZ — GIS DATA VALIDATION DASHBOARD"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color=WHITE, bold=True, size=18)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A3:L3")
    ws["A3"] = (
        "Pre-ingestion QA/QC | Authoritative GIS reconciliation | "
        f"Delivery status: {summary['delivery_status']}"
    )
    ws["A3"].font = Font(color=DARK, italic=True, size=10)
    ws["A3"].alignment = Alignment(horizontal="center")

    _card(ws, 1, "RECEIVED", summary["records_received"], NAVY)
    _card(ws, 3, "READY", summary["ready_to_update"], GREEN)
    _card(ws, 5, "NO CHANGE", summary["no_change"], BLUE)
    _card(ws, 7, "REJECTED", summary["rejected"], RED)
    _card(ws, 9, "UNMATCHED", summary["unmatched"], AMBER)
    _card(ws, 11, "QUALITY %", summary["data_quality_score_pct"], GRAY)

    ws["A9"] = "Validation Status"
    ws["B9"] = "Count"
    status_counts = classified["validation_status"].value_counts()
    status_order = ["READY_TO_UPDATE", "NO_CHANGE", "REJECTED", "UNMATCHED"]
    for row_idx, status in enumerate(status_order, start=10):
        ws.cell(row_idx, 1, status)
        ws.cell(row_idx, 2, int(status_counts.get(status, 0)))

    ws["A9"].fill = PatternFill("solid", fgColor=NAVY)
    ws["B9"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A9"].font = Font(color=WHITE, bold=True)
    ws["B9"].font = Font(color=WHITE, bold=True)

    doughnut = DoughnutChart()
    doughnut.title = "Delivery Status Distribution"
    doughnut.add_data(Reference(ws, min_col=2, min_row=9, max_row=13), titles_from_data=True)
    doughnut.set_categories(Reference(ws, min_col=1, min_row=10, max_row=13))
    doughnut.height = 7.2
    doughnut.width = 10
    ws.add_chart(doughnut, "D9")

    rejected = reports.get("Rejected", pd.DataFrame())
    if not rejected.empty:
        reasons = rejected["validation_reason"].value_counts().rename_axis("Reason").reset_index(name="Count")
    else:
        reasons = pd.DataFrame({"Reason": ["None"], "Count": [0]})

    ws["A17"] = "Rejection Reason"
    ws["B17"] = "Count"
    for i, row in reasons.iterrows():
        ws.cell(18 + i, 1, str(row["Reason"]))
        ws.cell(18 + i, 2, int(row["Count"]))

    for cell in (ws["A17"], ws["B17"]):
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)

    bar = BarChart()
    bar.type = "bar"
    bar.title = "Rejected Records by Reason"
    bar.y_axis.title = "Reason"
    bar.x_axis.title = "Records"
    bar.add_data(Reference(ws, min_col=2, min_row=17, max_row=17 + len(reasons)), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=18, max_row=17 + len(reasons)))
    bar.height = 8
    bar.width = 13
    ws.add_chart(bar, "D24")

    ws["A38"] = "Field-level proposed changes"
    ws["B38"] = summary["field_changes"]
    ws["A39"] = "Processing time (sec)"
    ws["B39"] = summary["processing_seconds"]
    ws["A40"] = "Delivery status"
    ws["B40"] = summary["delivery_status"]
    ws["A42"] = "Interpretation"
    ws["B42"] = (
        "READY FOR GIS means no rejected or unmatched records. "
        "REVIEW REQUIRED means exceptions must be reviewed before production use."
    )
    ws["B42"].alignment = Alignment(wrap_text=True)

    for col in range(1, 13):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18

    summary_df = pd.DataFrame({"Metric": list(summary.keys()), "Value": list(summary.values())})
    all_sheets = {"Summary": summary_df, **reports}
    for sheet_name, df in all_sheets.items():
        out = wb.create_sheet(sheet_name[:31])
        _write_dataframe(out, df)

        if "validation_status" in df.columns:
            status_col = list(df.columns).index("validation_status") + 1
            for col_cells in out.iter_cols(min_col=status_col, max_col=status_col, min_row=2, max_row=out.max_row):
                for cell in col_cells:
                    color = STATUS_COLORS.get(cell.value)
                    if color:
                        cell.fill = PatternFill("solid", fgColor=color)
                        cell.font = Font(color=WHITE, bold=True)

    wb.save(output_path)
    return output_path
